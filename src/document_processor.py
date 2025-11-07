"""Module d'extraction de texte depuis différents formats de documents"""
import io
from typing import Dict
import fitz  # PyMuPDF


def extract_text_from_pdf(file) -> str:
    """
    Extrait le texte d'un fichier PDF
    
    Args:
        file: Objet fichier (file-like object avec méthode .read())
        
    Returns:
        str: Texte extrait du PDF
    """
    try:
        
        # Lire le PDF depuis le buffer
        pdf_bytes = file.read()
        pdf_document = fitz.open(stream=pdf_bytes, filetype="pdf")
        
        text_content = []
        
        for page_num in range(pdf_document.page_count):
            page = pdf_document[page_num]
            text_content.append(page.get_text())
        
        pdf_document.close()
        
        full_text = "\n\n".join(text_content)
        return full_text.strip()
    
    except Exception as e:
        raise Exception(f"Erreur lors de l'extraction du PDF: {str(e)}")


def extract_text_from_docx(file) -> str:
    """
    Extrait le texte d'un fichier DOCX
    
    Args:
        file: Objet fichier (file-like object avec méthode .read())
        
    Returns:
        str: Texte extrait du DOCX
    """
    try:
        from docx import Document
        
        # Lire le DOCX depuis le buffer
        doc = Document(io.BytesIO(file.read()))
        
        text_content = []
        
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_content.append(paragraph.text)
        
        # Extraire aussi le texte des tableaux
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells])
                if row_text.strip():
                    text_content.append(row_text)
        
        full_text = "\n\n".join(text_content)
        return full_text.strip()
    
    except Exception as e:
        raise Exception(f"Erreur lors de l'extraction du DOCX: {str(e)}")


def extract_text_from_txt(file) -> str:
    """
    Extrait le texte d'un fichier TXT
    
    Args:
        file: Objet fichier (file-like object avec méthode .read())
        
    Returns:
        str: Texte extrait du TXT
    """
    try:
        # Essayer UTF-8 en premier
        text = file.read().decode('utf-8')
        return text.strip()
    except UnicodeDecodeError:
        # Fallback sur latin-1
        file.seek(0)
        text = file.read().decode('latin-1')
        return text.strip()


def extract_text_from_file(uploaded_file) -> str:
    """
    Extrait le texte d'un fichier selon son type
    
    Args:
        uploaded_file: Objet fichier avec attribut .type et méthode .read()
        
    Returns:
        str: Texte extrait
    """
    file_type = uploaded_file.type
    
    if file_type == "application/pdf":
        return extract_text_from_pdf(uploaded_file)
    elif file_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return extract_text_from_docx(uploaded_file)
    elif file_type == "text/plain":
        return extract_text_from_txt(uploaded_file)
    else:
        raise ValueError(f"Type de fichier non supporté: {file_type}")


def format_documents_context(uploaded_docs: Dict[str, str]) -> str:
    """
    Formate les documents uploadés en contexte structuré avec balises XML
    
    Args:
        uploaded_docs: Dictionnaire {filename: content}
        
    Returns:
        str: Contexte formaté avec balises XML
    """
    if not uploaded_docs:
        return ""
    
    result = "\n\n## DOCUMENTS DE RÉFÉRENCE FOURNIS PAR LE CHEF DE PROJET\n\n"
    result += "Le chef de projet a fourni les documents suivants. Tu DOIS les utiliser comme référence pour répondre avec précision et éviter d'inventer des informations.\n\n"
    result += "<documents_reference>\n"
    
    for filename, content in uploaded_docs.items():
        result += f"<document>\n"
        result += f"<filename>{filename}</filename>\n"
        result += f"<content>\n{content}\n</content>\n"
        result += f"</document>\n\n"
    
    result += "</documents_reference>\n\n"
    result += "⚠️ IMPORTANT: Utilise UNIQUEMENT les informations contenues dans les balises <documents_reference> pour les questions concernant ce projet spécifique.\n\n"
    
    return result

